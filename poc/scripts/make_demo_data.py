"""Generate SIMULATED enrichment (step 11 of the data guide) and SYNTHETIC documents for the shared folder.

Usage: python scripts/make_demo_data.py   (run from poc/)
Deterministic (seed 42). Everything produced here is fake and labelled as such.
"""
import json
import random
import sys
from datetime import date
from pathlib import Path

import docx
import fitz
import numpy as np
import pandas as pd
from PIL import Image, ImageDraw, ImageFilter, ImageFont

POC = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(POC))
from data_tools import PO, SUBS, T, WEIGHTS, df  # noqa: E402

ENR = POC / "data" / "enrichment"
DOCS = POC / "sample_docs"
RAW = POC / "data" / "web_raw"
rng = np.random.default_rng(42)
random.seed(42)
TODAY = date(2026, 10, 18)
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_B = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
FONT_S = "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf"

CATS = {
    "critical": ["Bogies & wheelsets", "Brake systems", "Traction & propulsion", "Signalling & ETCS",
                 "Pantographs & HV equipment", "Couplers & draft gear"],
    "mid": ["HVAC", "Doors & access systems", "Passenger information & IT", "Electronics & PCBs",
            "Castings & forgings", "Track & infrastructure"],
    "simple": ["Interior & seating", "Fasteners & MRO consumables", "Cables & connectors", "Logistics services"],
}
CAT_MIX = {"Complex": (.70, .25, .05), "Moderate": (.25, .55, .20), "Simple": (.05, .30, .65)}
LEAD = {"Bogies & wheelsets": (30, 52), "Brake systems": (20, 40), "Traction & propulsion": (30, 60),
        "Signalling & ETCS": (26, 52), "Pantographs & HV equipment": (16, 30), "Couplers & draft gear": (16, 30),
        "HVAC": (12, 26), "Doors & access systems": (12, 24), "Passenger information & IT": (8, 20),
        "Electronics & PCBs": (12, 52), "Castings & forgings": (20, 40), "Track & infrastructure": (10, 30),
        "Interior & seating": (6, 16), "Fasteners & MRO consumables": (1, 4), "Cables & connectors": (4, 12),
        "Logistics services": (1, 2)}
WELDING = {"Bogies & wheelsets", "Castings & forgings", "Couplers & draft gear", "Track & infrastructure"}
COUNTRIES = {
    "low": [("DE", "Germany"), ("FR", "France"), ("AT", "Austria"), ("CH", "Switzerland"), ("NL", "Netherlands"),
            ("SE", "Sweden"), ("DK", "Denmark"), ("BE", "Belgium")],
    "mid": [("PL", "Poland"), ("CZ", "Czechia"), ("ES", "Spain"), ("IT", "Italy"), ("PT", "Portugal"),
            ("HU", "Hungary"), ("RO", "Romania"), ("SK", "Slovakia")],
    "high": [("TR", "Turkey"), ("IN", "India"), ("CN", "China"), ("MA", "Morocco"), ("RS", "Serbia"), ("MX", "Mexico")],
}
FORCED = {  # keep the demo story consistent with the documents below
    "SUPP_09733": ("Brake systems", "ES", "Spain"),
    "SUPP_06901": ("Bogies & wheelsets", "PL", "Poland"),
    "SUPP_01319": ("Signalling & ETCS", "DE", "Germany"),
    "SUPP_05160": ("Electronics & PCBs", "CZ", "Czechia"),
    "SUPP_04684": ("Traction & propulsion", "IT", "Italy"),
    "SUPP_08780": ("Castings & forgings", "TR", "Turkey"),
    "SUPP_04226": ("Couplers & draft gear", "IN", "India"),
    "SUPP_04288": ("Doors & access systems", "PL", "Poland"),
}
FOCUS = list(FORCED)


# ---------------------------------------------------------------- enrichment
def build_enrichment(d: pd.DataFrame) -> pd.DataFrame:
    geo = d[["Environmental_Exposure_Score", "Geopolitical_Risk_Score", "Political_Stability_Score"]].mean(axis=1)
    lo, hi = geo.quantile([.33, .66])
    rows = []
    for (_, r), g in zip(d.iterrows(), geo):
        tier = rng.choice(["critical", "mid", "simple"], p=CAT_MIX[r["Supply Chain Complexity"]])
        cat = rng.choice(CATS[tier])
        band = "low" if g < lo else "mid" if g < hi else "high"
        cc, country = COUNTRIES[band][rng.integers(len(COUNTRIES[band]))]
        if r.supplierID in FORCED:
            cat, cc, country = FORCED[r.supplierID]
        q, esg = r.Quality_Risk_Score, r.ESG_Performance_Score
        iris = cat in CATS["critical"] + CATS["mid"] and rng.random() < (0.95 if q < 50 else 0.75 if q < 75 else 0.45)
        en15085 = cat in WELDING and rng.random() < (0.9 if q < 75 else 0.6)
        exp = lambda: date.fromordinal(TODAY.toordinal() + int(rng.integers(-60, 900))).isoformat()
        lt = LEAD[cat]
        rows.append({
            "supplierID": r.supplierID,
            "parent_company": " ".join(r["Supplier Name"].split()[:2]) + " Holding",
            "country_code": cc, "country": country, "geo_band": band, "category": cat,
            "category_criticality": tier,
            "iris_certified": bool(iris), "iris_expiry": exp() if iris else "",
            "en15085_certified": bool(en15085), "en15085_expiry": exp() if en15085 else "",
            "iso14001_certified": bool(esg < 60 and rng.random() < .8),
            "lead_time_weeks": int(rng.integers(lt[0], lt[1] + 1)),
            "safety_stock_weeks": int(rng.integers(2, 13)),
            "share_of_supplier_revenue_pct": round(float(np.clip(rng.lognormal(1.6, .7), .5, 60)), 1),
            "qualified_alternatives": int({"High": rng.integers(0, 2), "Medium": rng.integers(2, 4),
                                           "Low": rng.integers(4, 9)}[r["Sourcing Risk Level"]]),
            "data_label": "SIMULATED",
        })
    e = pd.DataFrame(rows)
    hero = e.supplierID == "SUPP_09733"
    e.loc[hero, ["iris_certified", "iris_expiry", "lead_time_weeks", "safety_stock_weeks", "qualified_alternatives"]] = \
        [True, "2026-12-31", 34, 4, 0]
    bog = e.supplierID == "SUPP_06901"
    e.loc[bog, ["en15085_certified", "en15085_expiry", "lead_time_weeks", "qualified_alternatives"]] = \
        [True, "2026-09-30", 38, 0]
    return e


def build_history(d: pd.DataFrame, focus_ids: set) -> pd.DataFrame:
    months = pd.period_range(end=pd.Period(TODAY, "M"), periods=24, freq="M")
    out = []
    sub = d[d.supplierID.isin(focus_ids)]
    for _, r in sub.iterrows():
        trend = 1.6 if r.supplierID in ("SUPP_09733", "SUPP_04288", "SUPP_04684") else rng.normal(0, .25)
        series = {}
        for c in SUBS:
            e = np.zeros(24)
            for t in range(1, 24):
                e[t] = .85 * e[t - 1] + rng.normal(0, 3.0)
            e = e - e[-1] * np.linspace(0, 1, 24)
            drift = trend * (np.arange(24) - 23) if c in ("Financial_Health_Score", "Credit_Rating_Score",
                                                           "Operational_Resilience_Score", "Quality_Risk_Score") else 0
            series[c] = np.clip(np.round(r[c] + e + drift), 0, 100)
        total = np.round(sum(WEIGHTS[c] * series[c] for c in SUBS))
        ops = (series["Operational_Resilience_Score"] + series["Quality_Risk_Score"]) / 2
        ops_now = (r.Operational_Resilience_Score + r.Quality_Risk_Score) / 2
        deliv = np.clip(r["Supplier Delivery Rating"] - .3 * (ops - ops_now) + rng.normal(0, 1.2, 24), 60, 100)
        deliv[-1] = r["Supplier Delivery Rating"]
        for i, m in enumerate(months):
            out.append({"supplierID": r.supplierID, "month": str(m), "total_risk": int(total[i]),
                        "financial_health": int(series["Financial_Health_Score"][i]),
                        "quality_risk": int(series["Quality_Risk_Score"][i]),
                        "operational_resilience": int(series["Operational_Resilience_Score"][i]),
                        "geopolitical": int(series["Geopolitical_Risk_Score"][i]),
                        "delivery_rating": round(float(deliv[i]), 1), "data_label": "SIMULATED"})
    return pd.DataFrame(out)


FLEETS = [
    ("RegioFlex EMU", 120, ["Bogies & wheelsets", "Brake systems", "Traction & propulsion", "Doors & access systems",
                            "HVAC", "Passenger information & IT", "Couplers & draft gear", "Interior & seating"]),
    ("ICX-300 high-speed", 60, ["Bogies & wheelsets", "Brake systems", "Traction & propulsion", "Signalling & ETCS",
                                "Pantographs & HV equipment", "HVAC", "Electronics & PCBs", "Doors & access systems"]),
    ("CargoMax freight loco", 80, ["Bogies & wheelsets", "Brake systems", "Traction & propulsion", "Couplers & draft gear",
                                   "Signalling & ETCS", "Castings & forgings", "Cables & connectors"]),
]


def build_bom(d: pd.DataFrame, e: pd.DataFrame) -> pd.DataFrame:
    m = d[["supplierID", PO, T, "Sourcing Risk Level"]].merge(e[["supplierID", "category"]], on="supplierID")
    rows = []
    for fleet, units, cats in FLEETS:
        for cat in cats:
            pool = m[m.category == cat].sort_values(PO, ascending=False).head(40)
            forced = [s for s, (c, *_) in FORCED.items() if c == cat]
            primary = forced[0] if forced and fleet != "CargoMax freight loco" else pool.sample(1, random_state=int(rng.integers(1e6))).supplierID.iloc[0]
            src = m.loc[m.supplierID == primary, "Sourcing Risk Level"].iloc[0]
            second = "" if src == "High" else pool[pool.supplierID != primary].sample(1, random_state=1).supplierID.iloc[0]
            rows.append({"fleet": fleet, "trainsets": units, "assembly": cat, "primary_supplier": primary,
                         "second_source": second, "single_source": src == "High",
                         "units_per_trainset": int(rng.integers(1, 9)), "data_label": "SIMULATED"})
    return pd.DataFrame(rows)


# ---------------------------------------------------------------- documents
def scan_effect(img: Image.Image, angle: float = .6) -> Image.Image:
    img = img.convert("L").rotate(angle, expand=False, fillcolor=245, resample=Image.BICUBIC)
    arr = np.asarray(img).astype(np.int16) + rng.normal(0, 9, (img.height, img.width)).astype(np.int16)
    return Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8)).filter(ImageFilter.GaussianBlur(.6))


def page(w=1240, h=1754) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    img = Image.new("RGB", (w, h), "white")
    return img, ImageDraw.Draw(img)


def f(size, bold=False, serif=False):
    return ImageFont.truetype(FONT_S if serif else FONT_B if bold else FONT, size)


def certificate_png(path, title, standard, holder, sid, scope, number, issued, expiry, body, extra=""):
    img, dr = page()
    dr.rectangle([40, 40, 1200, 1714], outline="black", width=6)
    dr.rectangle([60, 60, 1180, 1694], outline="black", width=2)
    dr.text((620, 170), "CERTIFICATE", font=f(78, serif=True), fill="black", anchor="mm")
    dr.text((620, 260), title, font=f(34, serif=True), fill="black", anchor="mm")
    dr.text((620, 330), f"according to {standard}", font=f(28), fill="black", anchor="mm")
    y = 450
    for k, v in [("Certificate holder", holder), ("Supplier ID", sid), ("Scope", scope),
                 ("Certificate No.", number), ("Date of issue", issued), ("Valid until", expiry),
                 ("Certification body", body)]:
        dr.text((140, y), k + ":", font=f(28, bold=True), fill="black")
        dr.multiline_text((520, y), v, font=f(28), fill="black", spacing=8)
        y += 90 + 38 * v.count("\n")
    if extra:
        dr.multiline_text((140, y + 20), extra, font=f(24), fill="black", spacing=8)
    dr.ellipse([860, 1380, 1100, 1620], outline="#1d3f8f", width=6)
    dr.text((980, 1500), "CERTIFIED\nSEAL", font=f(26, bold=True), fill="#1d3f8f", anchor="mm", align="center")
    dr.line([140, 1560, 560, 1560], fill="black", width=2)
    dr.text((140, 1575), "Lead auditor, signature", font=f(22), fill="black")
    dr.text((620, 1660), "SYNTHETIC DEMO DOCUMENT - NOT A REAL CERTIFICATE", font=f(20, bold=True), fill="#b00", anchor="mm")
    scan_effect(img, random.uniform(-.9, .9)).save(path, quality=88)


def table_scan(title, header, rows, footer):
    img, dr = page()
    dr.text((90, 90), title, font=f(40, bold=True), fill="black")
    dr.text((90, 150), "SYNTHETIC DEMO DOCUMENT", font=f(22, bold=True), fill="#b00")
    x0, y, cw = 90, 230, (1240 - 180) // len(header)
    for i, h in enumerate(header):
        dr.text((x0 + i * cw + 8, y + 10), h, font=f(24, bold=True), fill="black")
    dr.rectangle([x0, y, 1150, y + 50], outline="black", width=2)
    for r in rows:
        y += 50
        dr.rectangle([x0, y, 1150, y + 50], outline="black", width=1)
        for i, v in enumerate(r):
            dr.text((x0 + i * cw + 8, y + 12), str(v), font=f(23), fill="black")
    dr.multiline_text((90, y + 90), footer, font=f(24), fill="black", spacing=10)
    return scan_effect(img, random.uniform(-.7, .7))


def text_pdf(path, title, body):
    doc = fitz.open()
    pg = doc.new_page(width=595, height=842)
    pg.insert_textbox(fitz.Rect(50, 50, 545, 100), title, fontsize=16, fontname="hebo")
    pg.insert_textbox(fitz.Rect(50, 95, 545, 112), "SYNTHETIC DEMO DOCUMENT", fontsize=9, fontname="hebo", color=(.7, 0, 0))
    rc = pg.insert_textbox(fitz.Rect(50, 120, 545, 800), body, fontsize=10.5, fontname="helv", lineheight=1.35)
    if rc < 0:
        raise ValueError(f"text overflow in {path}")
    doc.save(path)


def scanned_pdf(path, images):
    doc = fitz.open()
    for im in images:
        tmp = ENR / "_tmp_scan.jpg"
        im.save(tmp, quality=85)
        pg = doc.new_page(width=595, height=842)
        pg.insert_image(pg.rect, filename=str(tmp))
        tmp.unlink()
    doc.save(path)


def build_documents(d, e, hist, bom):
    DOCS.mkdir(exist_ok=True)
    for sub in ["certificates", "audits", "quality", "contracts", "finance", "supplier_master", "market", "background_public"]:
        (DOCS / sub).mkdir(exist_ok=True)
    name = dict(zip(d.supplierID, d["Supplier Name"]))
    E = e.set_index("supplierID")

    # certificates (images -> OCR)
    certificate_png(DOCS / "certificates/IRIS_cert_SUPP_09733_scan.jpg", "Railway Quality Management System",
                    "ISO/TS 22163:2023 (IRIS Certification)", f"{name['SUPP_09733']}\nPlant Valencia (ES)", "SUPP_09733",
                    "Design and manufacture of brake control units\nand brake valve assemblies for rolling stock",
                    "IRIS-2023-ES-04417", "2024-01-08", "2026-12-31", "RailCert Europe GmbH (fictitious)",
                    "Note: surveillance audit 2026 raised 2 MAJOR findings.\nRecertification audit not yet scheduled.")
    certificate_png(DOCS / "certificates/EN15085_cert_SUPP_06901_scan.jpg", "Welding of railway vehicles and components",
                    "EN 15085-2, certification level CL1", f"{name['SUPP_06901']}\nPlant Wroclaw (PL)", "SUPP_06901",
                    "Welding of bogie frames and bolsters", "EN15085-PL-2021-0932", "2021-10-01", "2026-09-30",
                    "Instytut Spawalnictwa Cert (fictitious)",
                    "Status at 2026-10-18: EXPIRED - renewal audit pending.")
    certificate_png(DOCS / "certificates/IRIS_cert_SUPP_01319_scan.jpg", "Railway Quality Management System",
                    "ISO/TS 22163:2023 (IRIS Certification)", f"{name['SUPP_01319']}\nPlant Braunschweig (DE)", "SUPP_01319",
                    "Development and production of ETCS onboard units", "IRIS-2025-DE-11873", "2025-03-14", "2028-03-13",
                    "RailCert Europe GmbH (fictitious)")

    # delivery performance table scan for the hero supplier (from SIMULATED history)
    h = hist[hist.supplierID == "SUPP_09733"].tail(12)
    rows = [(r.month, f"{r.delivery_rating:.1f}%", r.total_risk, r.financial_health) for r in h.itertuples()]
    img = table_scan(f"Supplier scorecard - {name['SUPP_09733']} (SUPP_09733)",
                     ["Month", "On-time delivery", "Total risk", "Financial health"], rows,
                     "Category: Brake systems (RegioFlex EMU, ICX-300)\nTrend: on-time delivery falling, financial risk rising.\n"
                     "Prepared by: Supplier Quality, Q3 2026 business review")
    img.save(DOCS / "quality/scorecard_SUPP_09733_scan.png")

    # quality incident report: image-only PDF (forces OCR)
    r = d[d.supplierID == "SUPP_04288"].iloc[0]
    img1 = table_scan("Non-Conformance Report NCR-2026-0871",
                      ["Field", "Value", "", ""],
                      [("Supplier", "SUPP_04288", "", ""), ("Name", name["SUPP_04288"], "", ""),
                       ("Part", "Door drive unit DDU-4", "", ""), ("Fleet", "RegioFlex EMU", "", ""),
                       ("Defect", "Motor overheating", "", ""), ("Units affected", "37 of 220", "", ""),
                       ("Detected", "2026-09-22", "", ""), ("Severity", "Major", "", ""),
                       ("Delivery rating", f"{r['Supplier Delivery Rating']:.1f}% (expected ~91%)", "", "")],
                      "Root cause (supplier statement): change of motor sub-supplier in May 2026\nwithout notification (PCN missing).\n"
                      "Containment: 100% inspection at incoming goods.\nCorrective action due: 2026-11-15.")
    scanned_pdf(DOCS / "quality/NCR-2026-0871_SUPP_04288_scanned.pdf", [img1])

    # ESG audit (text PDF)
    text_pdf(DOCS / "audits/ESG_audit_SUPP_04226.pdf", f"Social & Environmental Audit - {name['SUPP_04226']} (SUPP_04226)",
             "Audit date: 2026-07-09 | Site: Pune (IN) | Category: Couplers & draft gear (SIMULATED)\n"
             "Framework: LkSG risk analysis / CSDDD-aligned on-site audit\n\n"
             "Summary: HIGH RISK. 3 critical and 4 major findings.\n\n"
             "Critical findings\n"
             "C1 Working hours: 72-hour weeks recorded in the forging shop for 11 of 13 weeks sampled.\n"
             "C2 Safety: two lost-time injuries at the drop forge in 2026; machine guards removed on press P-07.\n"
             "C3 Recruitment fees paid by migrant workers to labour agents (avg. 1.5 monthly wages).\n\n"
             "Major findings\n"
             "M1 No documented grievance mechanism accessible to contract workers.\n"
             "M2 Wastewater from pickling line discharged without treatment log.\n"
             "M3 Fire exits blocked in warehouse B.\n"
             "M4 No sub-supplier code of conduct for raw steel billets.\n\n"
             "Required actions (LkSG section 7 remedial plan): corrective action plan within 30 days, re-audit within 6 months, "
             "reimbursement of recruitment fees, restore machine guarding immediately.\n\n"
             "Auditor: Fictitious Assurance Ltd. This is a SYNTHETIC demo document.")

    # cyber incident notice (text PDF)
    text_pdf(DOCS / "audits/cyber_incident_notice_SUPP_05160.pdf", f"Security Incident Notification - {name['SUPP_05160']} (SUPP_05160)",
             "Date: 2026-10-02 | Category: Electronics & PCBs (SIMULATED) | Plant: Brno (CZ)\n\n"
             "On 2026-09-29 the supplier detected ransomware on its production planning servers (ERP/MES). "
             "Production of PCB assemblies was stopped for 6 days. Customer drawings for traction converter control boards "
             "may have been exfiltrated. Forensic investigation ongoing.\n\n"
             "Impact for us: 2 open purchase orders (EUR 1.2M) delayed by ~3 weeks. NIS2 supply-chain notification may be required.\n\n"
             "Requested actions: confirm affected drawings, rotate shared VPN credentials, request forensic report, "
             "review supplier against NIS2 supply-chain security requirements.\n\nSYNTHETIC demo document.")

    # finance memo (DOCX) - hidden financial risk
    r = d[d.supplierID == "SUPP_04684"].iloc[0]
    doc = docx.Document()
    doc.add_heading(f"Credit & Financial Alert - {name['SUPP_04684']} (SUPP_04684)", 1)
    doc.add_paragraph("SYNTHETIC DEMO DOCUMENT")
    for t in [f"Date: 2026-10-10. Category: Traction & propulsion (SIMULATED). Annual PO volume: EUR {r[PO] / 1e6:.1f}M.",
              f"Internal total risk index: {int(r[T])} (Medium) - looks acceptable, BUT Financial Health score is "
              f"{int(r.Financial_Health_Score)} (Critical).",
              "External signals: payment terms to sub-suppliers extended from 60 to 120 days; covenant waiver requested from "
              "lenders in August 2026; auditor's report includes a going-concern emphasis of matter.",
              "Dependency: we represent approx. 18% of the supplier's revenue (SIMULATED).",
              "Recommendation: request bank guarantee for prepayments, increase safety stock of traction motors from 4 to 10 "
              "weeks, start qualification of a second source (lead time 30+ weeks)."]:
        doc.add_paragraph(t)
    doc.save(DOCS / "finance/credit_alert_SUPP_04684.docx")

    # supplier self-assessment (markdown) - geopolitical hidden risk
    r = d[d.supplierID == "SUPP_08780"].iloc[0]
    (DOCS / "audits/self_assessment_SUPP_08780.md").write_text(
        f"# Supplier self-assessment questionnaire - {name['SUPP_08780']} (SUPP_08780)\n\nSYNTHETIC DEMO DOCUMENT\n\n"
        f"- Category: Castings & forgings (SIMULATED)\n- Main site: Kocaeli (TR); second site: none\n"
        f"- Internal total risk: {int(r[T])} (Low) but Geopolitical risk score: {int(r.Geopolitical_Risk_Score)} (Critical)\n"
        "- Export route: Black Sea / Bosporus shipping, 100% of volume\n- Business continuity plan: 'in preparation'\n"
        "- Currency exposure: costs in TRY, contracts in EUR; supplier asks for price escalation clause\n"
        "- Earthquake zone: site located in seismic zone 1; last structural assessment 2019\n")

    # contracts (keep the earlier ones, add one more)
    (DOCS / "contracts/framework_agreement_SUPP_09733_SYNTHETIC.md").write_text(
        f"# Framework agreement summary - {name['SUPP_09733']} (SUPP_09733)\n\nSYNTHETIC DEMO DOCUMENT\n\n"
        "| Item | Value |\n|---|---|\n| Scope | Brake control units BCU-7 for RegioFlex EMU and ICX-300 |\n"
        "| Term | 2024-01 to 2029-12 |\n| Annual volume | EUR 50.0M |\n| Single source | Yes |\n"
        "| Lead time | 34 weeks |\n| Safety stock agreed | 4 weeks |\n| Penalty | 0.3%/week of delay, cap 8% |\n"
        "| Exit | 18 months notice; tooling owned by us |\n| Escrow of drawings | No |\n")

    # supplier master + BOM + commodity (CSV)
    E.loc[FOCUS].reset_index().to_csv(DOCS / "supplier_master/supplier_master_focus_SIMULATED.csv", index=False)
    bom.to_csv(DOCS / "supplier_master/fleet_bom_SIMULATED.csv", index=False)
    months = pd.period_range(end=pd.Period(TODAY, "M"), periods=24, freq="M").astype(str)
    walk = lambda start, vol: np.round(start * np.exp(np.cumsum(rng.normal(0, vol, 24))), 1)
    pd.DataFrame({"month": months, "steel_hrc_index": walk(100, .045), "copper_index": walk(100, .04),
                  "aluminium_index": walk(100, .035), "semiconductor_leadtime_weeks": np.clip(np.round(14 + np.cumsum(rng.normal(0, 1.2, 24))), 8, 40),
                  "data_label": "SIMULATED"}).to_csv(DOCS / "market/commodity_indices_SIMULATED.csv", index=False)

    # public background (Wikipedia, CC BY-SA 4.0)
    for jf in sorted(RAW.glob("*.json")):
        pages = json.loads(jf.read_text()).get("query", {}).get("pages", {})
        for p in pages.values():
            txt = (p.get("extract") or "").strip()
            if len(txt) < 500:
                continue
            title = p["title"]
            url = "https://en.wikipedia.org/wiki/" + title.replace(" ", "_")
            (DOCS / "background_public" / f"wikipedia_{jf.stem}.md").write_text(
                f"# {title}\n\nSource: {url}\nLicense: CC BY-SA 4.0 (Wikipedia contributors). Excerpt, retrieved 2026-10-18.\n\n"
                + txt[:9000])

    # move earlier sample docs into folders
    for old, new in [("audit_report_SUPP_09733_SYNTHETIC.md", "audits"), ("contract_summary_SUPP_06901_SYNTHETIC.txt", "contracts")]:
        if (DOCS / old).exists():
            (DOCS / old).rename(DOCS / new / old)

    (DOCS / "_INDEX.md").write_text(
        "# Shared folder - demo documents\n\nEverything here is SYNTHETIC or SIMULATED (except background_public = Wikipedia, CC BY-SA).\n\n"
        "| Folder | What | Extraction |\n|---|---|---|\n"
        "| certificates/ | IRIS and EN 15085 certificates as scanned images | NVIDIA OCR |\n"
        "| quality/ | scorecard table scan (PNG), NCR as image-only PDF | NVIDIA OCR |\n"
        "| audits/ | ESG audit + cyber incident (text PDF), self-assessment (md), IRIS audit (md) | text |\n"
        "| finance/ | credit alert (DOCX) | text |\n| contracts/ | contract summaries | text |\n"
        "| supplier_master/ | country, category, certificates for focus suppliers; fleet BOM | text |\n"
        "| market/ | commodity indices | text |\n| background_public/ | rail disruptions, ETCS, bogies, brakes, CSDDD | text |\n\n"
        "Focus suppliers: " + ", ".join(FOCUS) + "\n")


def main():
    d = df()
    ENR.mkdir(parents=True, exist_ok=True)
    e = build_enrichment(d)
    e.to_csv(ENR / "suppliers_enriched_SIMULATED.csv", index=False)
    focus = set(d[d.priority].supplierID) | set(d[d.hidden_risk].nlargest(150, PO).supplierID) | set(FOCUS)
    hist = build_history(d, focus)
    hist.to_csv(ENR / "risk_history_monthly_SIMULATED.csv", index=False)
    bom = build_bom(d, e)
    bom.to_csv(ENR / "fleet_bom_SIMULATED.csv", index=False)
    build_documents(d, e, hist, bom)
    print(f"enrichment {len(e)} rows | history {len(hist)} rows ({hist.supplierID.nunique()} suppliers) | bom {len(bom)} rows")
    print(e.category.value_counts().to_string())


if __name__ == "__main__":
    main()
